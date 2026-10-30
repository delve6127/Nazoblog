"""sync_nazo_data.py가 nazo_data.json·검색어_맵핑.xlsx를 고쳐 쓰는 부분의 시험.

노션·네트워크·git은 부르지 않고, 실제 파일은 임시 폴더에 복사한 사본만 고친다.
실행 (저장소 루트에서): python3 -m unittest discover -s tests -v
"""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import sync_nazo_data as sync  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None

NEW_TITLE = "테스트 새 리뷰"
NEW_ENTRY = {"num": 9999, "date": "2026-10-30", "brand": "테스트", "rating": "추천",
             "quote": None, "beginner": False, "url": "/nazotoki-reviews/test"}


class SyncFileWriteTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.json_path = self.tmp / "nazo_data.json"
        shutil.copy(ROOT / "nazo_data.json", self.json_path)  # 실제 파일 형식 그대로, 사본에서만
        self.before = json.loads(self.json_path.read_text(encoding="utf-8"))

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def load(self):
        return json.loads(self.json_path.read_text(encoding="utf-8"))  # 파일이 깨지면 여기서 실패

    def add_new_entry(self):
        sort_data = self.load()["SORT_DATA"]
        sort_data[NEW_TITLE] = dict(NEW_ENTRY)
        sync.write_json_in_place(self.json_path, sort_data, updated_titles=[], new_titles=[NEW_TITLE])

    def test_update_and_add_keep_valid_json(self):
        sort_data = self.load()["SORT_DATA"]
        title = next(iter(sort_data))
        sort_data[title]["satisfaction"] = 1.5
        # 노션 한줄평에 들어갈 수 있는 따옴표·역슬래시·줄바꿈 (2026-10-30 이 시험으로 줄바꿈·역슬래시 깨짐 버그 발견)
        sort_data[title]["quote"] = '따옴표 "와 역슬래시 \\ 가 든\n두 줄 한줄평'
        sort_data[NEW_TITLE] = dict(NEW_ENTRY)

        sync.write_json_in_place(self.json_path, sort_data, updated_titles=[title], new_titles=[NEW_TITLE])

        after = self.load()
        self.assertEqual(after["SORT_DATA"][title]["satisfaction"], 1.5)
        self.assertEqual(after["SORT_DATA"][title]["quote"], sort_data[title]["quote"])
        self.assertEqual(after["SORT_DATA"][NEW_TITLE], NEW_ENTRY)
        self.assertEqual(len(after["SORT_DATA"]), len(self.before["SORT_DATA"]) + 1)
        others = {k: v for k, v in after["SORT_DATA"].items() if k not in (title, NEW_TITLE)}
        self.assertEqual(others, {k: v for k, v in self.before["SORT_DATA"].items() if k != title})
        self.assertEqual(after["PHONETIC_MAP"], self.before["PHONETIC_MAP"])
        self.assertEqual(after["COMPANY_MAP"], self.before["COMPANY_MAP"])

    @unittest.skipIf(openpyxl is None, "openpyxl 미설치")
    def test_xlsx_keywords_go_into_phonetic_map(self):
        self.add_new_entry()  # SORT_DATA에는 있고 PHONETIC_MAP에는 없는 리뷰
        xlsx = self.tmp / "검색어_맵핑.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["구분", "원본 이름 (키)", "검색 키워드들 (쉼표로 구분)"])
        ws.append(["나조 이름", NEW_TITLE, "테스트키워드, 테스트 두번째"])
        wb.save(xlsx)

        with mock.patch.object(sync, "XLSX_PATH", xlsx):
            result = sync.apply_xlsx_to_maps(self.json_path, self.load())

        after = self.load()
        self.assertEqual(after["PHONETIC_MAP"][NEW_TITLE.lower()], ["테스트키워드", "테스트 두번째"])
        self.assertEqual(len(result["phonetic"]), 1)
        self.assertEqual(after["COMPANY_MAP"], self.before["COMPANY_MAP"])

    @unittest.skipIf(openpyxl is None, "openpyxl 미설치")
    def test_update_xlsx_saves_only_when_rows_added(self):
        xlsx = self.tmp / "검색어_맵핑.xlsx"
        shutil.copy(ROOT / "검색어_맵핑.xlsx", xlsx)
        original = xlsx.read_bytes()

        with mock.patch.object(sync, "XLSX_PATH", xlsx):
            sync.update_xlsx([], [])
            self.assertEqual(xlsx.read_bytes(), original)  # 추가할 것 없으면 파일 그대로 (2026-10-30 수정)
            sync.update_xlsx([NEW_TITLE], [])

        names = [row[1] for row in openpyxl.load_workbook(xlsx).active.iter_rows(values_only=True)]
        self.assertIn(NEW_TITLE, names)


if __name__ == "__main__":
    unittest.main()
